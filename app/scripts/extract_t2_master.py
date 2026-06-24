#!/usr/bin/env python3
"""
=============================================================================
SCRIPT NAME: extract_t2_master.py
=============================================================================

INPUT FILES:
- /Users/arjundivecha/Dropbox/AAA Backup/A Complete/T2 Factor Timing Fuzzy/T2 Master.xlsx
  (default --input) Multi-sheet Excel workbook. Each sheet: column A = dates,
  row 1 = country headers, body = numeric factor values.

OUTPUT FILES:
- /Users/arjundivecha/Dropbox/AAA Backup/A Working/Triptych/app/data/t2_master.json
  (default --output) Columnar JSON consumed by the Triptych web frontend.
  Format v2: per sheet a "dates" array plus one value-array per country
  (null for missing). Written compactly (no indentation) and atomically
  (temp file + rename).

VERSION: 2.0
LAST UPDATED: 2026-06-10
AUTHOR: Arjun Divecha

DESCRIPTION:
Converts the T2 Master workbook into a chart-ready JSON dataset for the
Triptych web app. Reads every sheet of the workbook in read-only mode,
treats the first row as country names and column A as dates, and emits a
columnar structure:

  {
    "format": 2,
    "generated_at": "<UTC ISO timestamp of this extraction>",
    "source_file": "<absolute path of the input workbook>",
    "source_mtime": "<UTC ISO timestamp of the workbook's mtime>",
    "sheets": {
      "<sheet name>": {
        "countries": ["India", ...],
        "dates": ["2000-02-01", ...],
        "values": { "India": [21.3, null, ...], ... }
      }
    }
  }

Columnar layout stores each date and country name once instead of per-row,
which makes the file roughly 3x smaller and much faster to parse than the
old row-oriented format. Rows with no usable date and sheets with no data
are skipped.

DEPENDENCIES:
- openpyxl

USAGE:
python3 extract_t2_master.py            # uses the default input/output above
python3 extract_t2_master.py --input <path.xlsx> --output <path.json>
=============================================================================
"""

from __future__ import annotations

import argparse
import json
import os
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DEFAULT_INPUT = (
    "/Users/arjundivecha/Dropbox/AAA Backup/A Complete/"
    "T2 Factor Timing Fuzzy/T2 Master.xlsx"
)
DEFAULT_OUTPUT = (
    "/Users/arjundivecha/Dropbox/AAA Backup/A Working/"
    "Triptych/app/data/t2_master.json"
)


def normalize_date(value: Any) -> str | None:
    """Convert Excel date-like values to YYYY-MM-DD strings."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return None
    return str(value)


def to_float(value: Any) -> float | None:
    """Convert numeric values to float; return None for blanks/non-numeric."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def extract_workbook(input_xlsx: Path) -> dict[str, Any]:
    wb = load_workbook(input_xlsx, data_only=True, read_only=True)
    source_mtime = datetime.fromtimestamp(input_xlsx.stat().st_mtime, UTC)
    output: dict[str, Any] = {
        "format": 2,
        "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "source_file": str(input_xlsx),
        "source_mtime": source_mtime.isoformat(timespec="seconds"),
        "sheets": {},
    }

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row < 2 or ws.max_column < 2:
            continue

        row_iter = ws.iter_rows(
            min_row=1,
            max_row=ws.max_row,
            min_col=1,
            max_col=ws.max_column,
            values_only=True,
        )

        try:
            header_row = next(row_iter)
        except StopIteration:
            continue

        headers: list[str | None] = []
        for raw in header_row[1:]:
            if raw is None or str(raw).strip() == "":
                headers.append(None)
            else:
                headers.append(str(raw).strip())

        countries = [h for h in headers if h]
        if not countries:
            continue

        dates: list[str] = []
        values: dict[str, list[float | None]] = {c: [] for c in countries}

        for raw_row in row_iter:
            dt = normalize_date(raw_row[0] if raw_row else None)
            if not dt:
                continue

            row_vals: dict[str, float | None] = {}
            has_any_value = False
            for header, raw_value in zip(headers, raw_row[1:]):
                if not header:
                    continue
                num = to_float(raw_value)
                row_vals[header] = num
                if num is not None:
                    has_any_value = True

            if not has_any_value:
                continue

            dates.append(dt)
            for c in countries:
                values[c].append(row_vals.get(c))

        if dates:
            output["sheets"][sheet_name] = {
                "countries": countries,
                "dates": dates,
                "values": values,
            }

    wb.close()
    return output


def write_atomic(payload: dict[str, Any], output_json: Path) -> None:
    """Write JSON compactly via a temp file then rename (atomic)."""
    output_json.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = output_json.with_suffix(".json.tmp")
    tmp_path.write_text(
        json.dumps(payload, separators=(",", ":")), encoding="utf-8"
    )
    os.replace(tmp_path, output_json)


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract workbook to JSON")
    parser.add_argument("--input", default=DEFAULT_INPUT, help="Path to input .xlsx")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help="Path to output .json")
    args = parser.parse_args()

    input_xlsx = Path(args.input).expanduser().resolve()
    output_json = Path(args.output).expanduser().resolve()

    if not input_xlsx.exists():
        raise FileNotFoundError(f"Input workbook not found: {input_xlsx}")

    payload = extract_workbook(input_xlsx)
    write_atomic(payload, output_json)
    print(f"Wrote {output_json}")
    print(f"Sheets exported: {len(payload['sheets'])}")


if __name__ == "__main__":
    main()
